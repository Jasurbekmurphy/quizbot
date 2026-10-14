"""Testlarni Excel (.xlsx), Word (.docx) yoki matn (.txt) fayldan o'qish.

Excel: Savol | To'g'ri javob | Javob 2 | Javob 3 | ... | Izoh (ixtiyoriy)
Word / matn — quyidagi ko'rinishlardan biri:
  1) shablon:            2) harfli variantlar:       3) HEMIS ko'rinishi:
     ? Savol matni          1. Savol matni              Savol matni
     + To'g'ri javob        A) javob                    ====
     - Noto'g'ri            *B) to'g'ri javob           #To'g'ri javob
     - Noto'g'ri            C) javob                    ====
                            yoki oxirida «Javob: B»     Noto'g'ri
                                                        ++++
"""
import io
import re

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

MAX_OPTIONS = 10


def _clean(value) -> str:
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return " ".join(str(value if value is not None else "").split())


def _norm(text) -> str:
    return re.sub(r"[ʻʼ‘’`´]", "'", str(text or "")).strip().lower()


def _question(text, options, correct, explanation=""):
    return dict(text=text, options=options, correct=correct, explanation=explanation)


def _check(q, where, errors):
    """Savolni tekshiradi. To'g'ri bo'lsa True."""
    if not q["text"]:
        errors.append(f"{where}: savol matni yo'q")
    elif len(q["options"]) < 2:
        errors.append(f"{where}: kamida 2 ta javob bo'lishi kerak")
    elif len(q["options"]) > MAX_OPTIONS:
        errors.append(f"{where}: javoblar {MAX_OPTIONS} tadan ko'p bo'lmasin")
    elif q["correct"] is None:
        errors.append(f"{where}: to'g'ri javob belgilanmagan")
    elif len(set(o.lower() for o in q["options"])) != len(q["options"]):
        errors.append(f"{where}: bir xil javoblar bor")
    else:
        return True
    return False


# ---------- Excel ----------

def parse_xlsx(data: bytes):
    wb = load_workbook(io.BytesIO(data), data_only=True, read_only=True)
    rows = [list(r) for r in wb.worksheets[0].iter_rows(values_only=True)]
    start, expl_col = 0, None
    if rows and any("savol" in _norm(c) or "вопрос" in _norm(c) for c in rows[0]):
        start = 1  # sarlavha qatori
        heads = [_norm(c) for c in rows[0]]
        expl_col = next((j for j, h in enumerate(heads) if "izoh" in h or "изох" in h or "пояснен" in h), None)
    questions, errors = [], []
    for n, row in enumerate(rows[start:], start=start + 1):
        cells = [_clean(c) for c in row]
        if not any(cells):
            continue
        text = cells[0]
        options = [c for j, c in enumerate(cells[1:], start=1) if c and j != expl_col]
        explanation = cells[expl_col] if expl_col is not None and expl_col < len(cells) else ""
        q = _question(text, options, 0 if options else None, explanation)
        if _check(q, f"{n}-qator", errors):
            questions.append(q)
    return questions, errors


# ---------- Word / matn ----------

def _docx_lines(data: bytes):
    import docx

    doc = docx.Document(io.BytesIO(data))
    lines = []
    for p in doc.paragraphs:
        lines += p.text.split("\n")
    for table in doc.tables:  # jadval ichida yozilgan testlar ham o'qiladi
        for row in table.rows:
            for cell in row.cells:
                lines += cell.text.split("\n")
    return lines


Q_NUM = re.compile(r"^\s*(?:\d{1,4}\s*[.)]|savol\s*\d+\s*[.:)]?|savol\s*[.:)]|\?)\s*", re.I)
OPT_MARK = re.compile(r"^\s*([+\-*#])\s*")
OPT_LETTER = re.compile(r"^\s*([*+#]?)\s*([A-Ja-jА-Да-д])\s*[).]\s*(.*)$")
ANSWER_LINE = re.compile(r"^\s*(?:to'g'ri\s+)?javob\s*[:\-]\s*([A-Ja-jА-Да-д])\b", re.I)
EXPL_LINE = re.compile(r"^\s*izoh\s*[:\-]\s*(.*)$", re.I)
LETTERS = "abcdefghij"
CYR_LETTERS = "абвгд"


def _letter_index(ch: str) -> int:
    ch = ch.lower()
    return LETTERS.index(ch) if ch in LETTERS else CYR_LETTERS.index(ch)


def _parse_hemis(text: str):
    questions, errors = [], []
    blocks = [b for b in re.split(r"^\s*\+{4,}\s*$", text, flags=re.M) if b.strip()]
    for n, block in enumerate(blocks, start=1):
        parts = [" ".join(p.split()) for p in re.split(r"^\s*={4,}\s*$", block, flags=re.M)]
        parts = [p for p in parts if p]
        if not parts:
            continue
        options, correct = [], None
        for p in parts[1:]:
            if p.startswith("#"):
                correct = len(options)
                p = p[1:].strip()
            options.append(p)
        q = _question(Q_NUM.sub("", parts[0], count=1).strip(), options, correct)
        if _check(q, f"{n}-savol", errors):
            questions.append(q)
    return questions, errors


def parse_text(text: str):
    text = text.replace("\r", "")
    if re.search(r"^\s*={4,}\s*$", text, flags=re.M):
        return _parse_hemis(text)
    questions, errors = [], []
    cur = None

    def finish():
        # Javobi yo'q qatorlar (sarlavha, yo'riqnoma) e'tiborsiz qoldiriladi
        if cur and cur["options"]:
            where = f"{len(questions) + len(errors) + 1}-savol («{cur['text'][:40]}»)"
            if _check(cur, where, errors):
                questions.append(cur)

    for raw in text.split("\n"):
        line = " ".join(raw.split())
        if not line:
            continue
        m_ans = ANSWER_LINE.match(_norm(line))
        m_expl = EXPL_LINE.match(line)
        m_mark = OPT_MARK.match(line)
        m_letter = OPT_LETTER.match(line)
        if cur and m_ans:  # «Javob: B»
            i = _letter_index(m_ans[1])
            cur["correct"] = i if i < len(cur["options"]) else None
        elif cur and m_expl:
            cur["explanation"] = m_expl[1]
        elif cur and m_letter:  # «A) javob», «*B) to'g'ri javob»
            if m_letter[1]:
                cur["correct"] = len(cur["options"])
            cur["options"].append(m_letter[3].strip())
        elif cur and m_mark:  # «+ to'g'ri», «- noto'g'ri»
            if m_mark[1] in "+*#":
                cur["correct"] = len(cur["options"])
            cur["options"].append(OPT_MARK.sub("", line, count=1).strip())
        elif Q_NUM.match(line) or cur is None or cur["options"]:
            finish()
            cur = _question(Q_NUM.sub("", line, count=1).strip(), [], None)
        else:
            cur["text"] += " " + line  # savol bir necha qatorda
    finish()
    return questions, errors


def parse_file(filename: str, data: bytes):
    """Natija: (savollar, xatolar). Har savol: {text, options, correct, explanation}."""
    name = filename.lower()
    if name.endswith(".xlsx"):
        return parse_xlsx(data)
    if name.endswith(".docx"):
        return parse_text("\n".join(_docx_lines(data)))
    if name.endswith(".txt"):
        return parse_text(data.decode("utf-8-sig", errors="replace"))
    return [], ["Faqat Excel (.xlsx), Word (.docx) yoki matn (.txt) fayl qabul qilinadi."]


# ---------- Shablonlar ----------

EXAMPLES = [
    ("O'zbekiston poytaxti qaysi shahar?", ["Toshkent", "Samarqand", "Buxoro", "Andijon"], "Toshkent 1930-yildan poytaxt."),
    ("5 × 6 = ?", ["30", "11", "25", "36"], ""),
    ("Suvning kimyoviy formulasi qaysi?", ["H₂O", "CO₂", "O₂", "NaCl"], ""),
]


def make_xlsx_template() -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Test"
    ws.append(["Savol", "To'g'ri javob", "Javob 2", "Javob 3", "Javob 4", "Izoh (ixtiyoriy)"])
    for q, opts, expl in EXAMPLES:
        ws.append([q] + opts + [expl])
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="2F6FB0")
        c.alignment = Alignment(horizontal="center", wrap_text=True)
    ws["B1"].fill = PatternFill("solid", fgColor="2E8B57")
    for col, width in zip("ABCDEF", (50, 22, 22, 22, 22, 30)):
        ws.column_dimensions[col].width = width
    ws.freeze_panes = "A2"
    note = wb.create_sheet("Yo'riqnoma")
    for line in [
        "1. Har bir qator — bitta savol.",
        "2. A ustun — savol, B ustun — TO'G'RI javob (yashil ustun).",
        "3. C, D, E ... ustunlar — noto'g'ri javoblar (2 tadan 10 tagacha javob bo'lishi mumkin).",
        "4. Bot javoblarni o'zi aralashtiradi, to'g'ri javob har doim B ustunda tursin.",
        "5. «Izoh» ustuni ixtiyoriy: javob berilgandan keyin ko'rsatiladi.",
        "6. Birinchi qatordagi sarlavhalarni o'chirmang.",
    ]:
        note.append([line])
    note.column_dimensions["A"].width = 90
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def make_docx_template() -> bytes:
    import docx

    doc = docx.Document()
    doc.add_heading("Test shabloni", level=1)
    doc.add_paragraph(
        "Har bir savol «?» belgisi bilan boshlanadi. To'g'ri javob oldiga «+», noto'g'ri javoblar oldiga «-» "
        "qo'yiladi. Javoblar 2 tadan 10 tagacha. Izoh (ixtiyoriy) «Izoh:» bilan yoziladi. "
        "Bot javoblarni o'zi aralashtiradi."
    )
    for q, opts, expl in EXAMPLES:
        doc.add_paragraph(f"? {q}")
        doc.add_paragraph(f"+ {opts[0]}")
        for o in opts[1:]:
            doc.add_paragraph(f"- {o}")
        if expl:
            doc.add_paragraph(f"Izoh: {expl}")
        doc.add_paragraph("")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
